"""Render extracted ScopeItems to a 2-sheet Excel workbook.

Sheet 1 — Scope Items
    One row per scope item. Columns:
      Scope Item ID | Scope Item Name | Process Description | Dependencies

Sheet 2 — Process Steps
    One row per process step (group-heading rows from the source table
    are excluded — every row is a real, actionable step). Columns:
      Scope Item ID | Scope Item Name | Process Step | Business Roles |
      Transaction/App | Expected Results

Conventions (per the xlsx skill guidance):
- Calibri font, bold header row, frozen header, auto-filter
- No formulas (this is a data export — no calculation risk)
- Wrapped text on long-text columns
- Newlines inside cells are preserved (Excel renders them with wrap-text on)
"""
from __future__ import annotations

import io
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from docx_parser import ScopeItem, Dependency, ProcessStep


# --- Styling constants ----------------------------------------------------
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", start_color="0070F2")  # SAP-ish blue
BODY_FONT = Font(name="Calibri", size=11)

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP_LEFT = Alignment(vertical="top")


# --- Helpers --------------------------------------------------------------

def _format_dependencies(deps: list[Dependency]) -> str:
    """Render the dependency list into a single human-readable cell.

    Format: one dep per line, '<ID> – <Name> (<kind>)'. When a dep has no
    name, the dash is omitted. Newlines in cells display as line breaks
    when the cell uses wrap-text.
    """
    if not deps:
        return ""
    lines: list[str] = []
    for d in sorted(deps, key=lambda x: x.scope_id):
        name_part = f" – {d.name}" if d.name else ""
        lines.append(f"{d.scope_id}{name_part} ({d.kind})")
    return "\n".join(lines)


def _normalize_cell_text(text: str | None) -> str:
    """Normalize a string for an Excel cell — trim and strip CR characters."""
    if not text:
        return ""
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def _apply_header_style(ws, num_cols: int) -> None:
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            wrap_text=True, vertical="center", horizontal="left",
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(num_cols)}{ws.max_row}"
    )
    ws.row_dimensions[1].height = 22


def _apply_column_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _apply_body_alignment(ws, wrap_columns: set[int]) -> None:
    """Set vertical-top alignment on all body cells; wrap-text on long columns."""
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP if cell.column in wrap_columns else TOP_LEFT


# --- Sheet builders -------------------------------------------------------

def _build_scope_items_sheet(ws, items: list[ScopeItem]) -> None:
    headers = [
        "Scope Item ID",
        "Scope Item Name",
        "Process Description",
        "Dependencies",
    ]
    ws.append(headers)

    for item in items:
        ws.append([
            item.scope_id,
            item.name or "",
            _normalize_cell_text(item.process_description),
            _format_dependencies(item.dependencies),
        ])

    _apply_header_style(ws, num_cols=len(headers))
    _apply_column_widths(ws, [14, 50, 80, 40])
    _apply_body_alignment(ws, wrap_columns={2, 3, 4})


def _build_process_steps_sheet(ws, items: list[ScopeItem]) -> None:
    headers = [
        "Scope Item ID",
        "Scope Item Name",
        "Process Step",
        "Business Roles",
        "Transaction/App",
        "Expected Results",
    ]
    ws.append(headers)

    for item in items:
        for step in item.process_steps:
            if step.group_heading:
                # Skip — these are visual headers in the source doc, not
                # actionable steps. Including them would force every column
                # except 'Process Step' to be empty and break filtering.
                continue
            ws.append([
                item.scope_id,
                item.name or "",
                _normalize_cell_text(step.process_step),
                _normalize_cell_text(step.business_role),
                _normalize_cell_text(step.transaction_app),
                _normalize_cell_text(step.expected_results),
            ])

    _apply_header_style(ws, num_cols=len(headers))
    _apply_column_widths(ws, [14, 35, 50, 30, 40, 50])
    _apply_body_alignment(ws, wrap_columns={2, 3, 4, 5, 6})


# --- Public API -----------------------------------------------------------

def render_xlsx(items: Iterable[ScopeItem]) -> bytes:
    """Build an in-memory .xlsx workbook and return its bytes.

    The returned bytes can be passed directly to st.download_button() —
    no temporary file needed.
    """
    items_list = sorted(items, key=lambda x: x.scope_id)

    wb = Workbook()
    # Replace the default sheet with our first sheet
    ws1 = wb.active
    ws1.title = "Scope Items"
    _build_scope_items_sheet(ws1, items_list)

    ws2 = wb.create_sheet("Process Steps")
    _build_process_steps_sheet(ws2, items_list)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
